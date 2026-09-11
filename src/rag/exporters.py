"""The datasheet as .xlsx and the script as .txt."""

from __future__ import annotations

import io
import re
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .schema import DATASHEET_COLUMNS, TestCase

SHEET_NAME = "Datasheet"

# Widths from the delivered datasheets: a wide wrapped Description with
# narrow fixed columns either side. Keyed by attribute, so reordering
# DATASHEET_COLUMNS cannot misalign them.
_COLUMN_WIDTHS = {
    "s_no": 6,
    "requirement": 13,
    "description": 105,
    "folder": 20,
    "optimization_technique": 22,
    "test_type": 11,
    "test_technique": 24,
    "retired": 9,
    "scorable": 10,
    "comments": 14,
}
_CONFIDENCE_LABELS = [
    "Confidence",
    "Confidence_Retrieval",
    "Confidence_Grounding",
    "Similarity_To_Existing",
    "Closest_Existing_Case",
    "Needs_Review",
]

_HEADER_FONT = Font(bold=True)
_COMMENTS_FILL = PatternFill("solid", fgColor="FFC000")


def datasheet_to_xlsx(test_cases: list[TestCase], include_confidence: bool = True) -> bytes:
    """The datasheet as an .xlsx byte string ready to stream to the browser.

    Columns A-J are exactly the delivered format, so the file drops into the
    existing test-execution flow unchanged. The confidence columns are
    appended after them (K onwards) rather than inserted, so anything
    reading the sheet by column position still works.
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME

    labels = [label for _, label in DATASHEET_COLUMNS]
    if include_confidence:
        labels += _CONFIDENCE_LABELS

    sheet.append(labels)
    for index, label in enumerate(labels, start=1):
        cell = sheet.cell(row=1, column=index)
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        if label == "Comments":
            cell.fill = _COMMENTS_FILL

    for test_case in test_cases:
        row = test_case.to_row()
        if include_confidence:
            confidence = test_case.confidence
            row += [
                confidence.overall,
                confidence.retrieval,
                confidence.grounding,
                confidence.similarity_to_existing,
                confidence.closest_existing_id or "",
                "Yes" if confidence.needs_review else "No",
            ]
        sheet.append(row)

    _apply_layout(sheet, len(labels))

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _apply_layout(sheet, column_count: int) -> None:
    for index, (attr, _label) in enumerate(DATASHEET_COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = _COLUMN_WIDTHS.get(attr, 16)
    for index in range(len(DATASHEET_COLUMNS) + 1, column_count + 1):
        sheet.column_dimensions[get_column_letter(index)].width = 14

    description_column = next(
        i for i, (attr, _) in enumerate(DATASHEET_COLUMNS, start=1) if attr == "description"
    )
    for row in sheet.iter_rows(min_row=2, max_col=column_count):
        for cell in row:
            cell.alignment = Alignment(
                vertical="top", wrap_text=cell.column == description_column
            )

    sheet.freeze_panes = "A2"


def script_to_text(script: str, requirement_id: str | None) -> bytes:
    """The automation script as a downloadable .txt with a provenance header,
    so a script that reaches a reviewer out of context cannot be mistaken
    for a hand-written one.
    """
    stamp = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    header = (
        f"# Generated test script draft for requirement "
        f"{requirement_id or '(not stated)'}\n"
        f"# Generated {stamp} - review before execution.\n"
    )
    body = script if script.endswith("\n") else script + "\n"
    # CRLF to match the existing Windows-authored scripts in the repository;
    # a mixed-ending diff would obscure real changes.
    return (header + body).replace("\r\n", "\n").replace("\n", "\r\n").encode("utf-8")


def download_name(kind: str, requirement_id: str | None, extension: str) -> str:
    """A filename a reviewer can identify without opening it."""
    stamp = datetime.now(timezone.utc).strftime("%Y%m%d-%H%M")
    safe_id = re.sub(r"[^A-Za-z0-9_.-]", "_", requirement_id or "requirement")
    return f"{safe_id}_{kind}_{stamp}.{extension}"
