"""One parser per source format, all behind the Extractor protocol.

PDF -> PyMuPDF, XLSX -> openpyxl, PY/PYI -> stdlib ast, HTML -> stdlib
html.parser, TXT -> plain text (or ast, when the .txt is really Python).
"""

from __future__ import annotations

import ast
import re
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path
from typing import Any, Protocol

import openpyxl

try:
    import pymupdf
except ImportError:  # PyMuPDF < 1.24.3 only ships the legacy `fitz` alias
    import fitz as pymupdf


@dataclass
class ExtractedUnit:
    """One logical unit of a source document, before chunking.

    A unit is the natural granularity of its format: a PDF section or table,
    a spreadsheet row, a Python function/class, an HTML track section.
    Chunking never mixes units of different unit_types.
    """

    text: str
    # section_text | table | test_case | function | class | module_docstring
    # | raw_python | text
    unit_type: str
    locator: dict[str, Any] = field(default_factory=dict)
    extra: dict[str, Any] = field(default_factory=dict)


class Extractor(Protocol):
    def extract(self, file_path: Path) -> list[ExtractedUnit]:
        ...


# --- PDF -------------------------------------------------------------------

_HEADING_SIZE_RATIO_DEFAULT = 1.15
_TABLE_REPEAT_THRESHOLD_DEFAULT = 0.5
_LINE_REPEAT_THRESHOLD = 0.5
_FALLBACK_HEADING_MAX_WORDS = 12
_MIN_TABLE_ROWS = 2
_BOLD_FLAG = 1 << 4  # PyMuPDF span flag bit for synthetic/real bold

_DIGITS_RE = re.compile(r"\d+")
_HEADING_NUMBER_RE = re.compile(r"^(\d+(?:\.\d+){0,5})\.?\s+\S")
_DOT_LEADER_RE = re.compile(r"\.{4,}")
_DECORATIVE_PREFIX_RE = re.compile(r"^[/\\|•·\s]+")


@dataclass
class Line:
    text: str
    top: float
    bottom: float
    size: float
    bold: bool
    page: int


@dataclass
class HeadingCandidate:
    level: int
    title: str


def group_words_into_lines(words: list[dict], page: int, tolerance: float = 2.0) -> list[Line]:
    """Clusters words into visual lines by vertical position, tolerant of the
    sub-pixel jitter between words on one line.
    """
    if not words:
        return []

    ordered = sorted(words, key=lambda w: (w["top"], w["x0"]))
    clusters: list[list[dict]] = []
    for word in ordered:
        if clusters and abs(word["top"] - clusters[-1][-1]["top"]) <= tolerance:
            clusters[-1].append(word)
        else:
            clusters.append([word])

    lines: list[Line] = []
    for cluster in clusters:
        cluster.sort(key=lambda w: w["x0"])
        lines.append(
            Line(
                text=" ".join(w["text"] for w in cluster),
                top=min(w["top"] for w in cluster),
                bottom=max(w["bottom"] for w in cluster),
                size=max(w.get("size", 0) for w in cluster),
                bold=any("bold" in str(w.get("fontname", "")).lower() for w in cluster),
                page=page,
            )
        )
    return lines


def compute_body_size(all_line_sizes: list[float]) -> float:
    """The most common line font size - the baseline headings are measured against."""
    if not all_line_sizes:
        return 0.0
    counts: dict[float, int] = {}
    for size in all_line_sizes:
        rounded = round(size, 1)
        counts[rounded] = counts.get(rounded, 0) + 1
    return max(counts.items(), key=lambda kv: kv[1])[0]


def build_size_rank(candidate_sizes: list[float]) -> dict[float, int]:
    """Distinct heading font sizes -> level, largest first, for unnumbered headings."""
    distinct = sorted({round(s, 1) for s in candidate_sizes}, reverse=True)
    return {size: min(i + 1, 4) for i, size in enumerate(distinct)}


class HeadingClassifier:
    def __init__(self, body_size: float, size_rank: dict[float, int], size_ratio: float = 1.15):
        self.body_size = body_size
        self.size_rank = size_rank
        self.size_ratio = size_ratio

    def classify(self, line: Line) -> HeadingCandidate | None:
        if _DOT_LEADER_RE.search(line.text):  # a table-of-contents line
            return None

        title = _DECORATIVE_PREFIX_RE.sub("", line.text).strip()
        if not title:
            return None

        match = _HEADING_NUMBER_RE.match(title)
        if match and (line.bold or line.size > self.body_size * 1.05):
            return HeadingCandidate(level=len(match.group(1).split(".")), title=title)

        is_large = self.body_size and line.size > self.body_size * self.size_ratio
        is_short = len(title.split()) <= _FALLBACK_HEADING_MAX_WORDS
        if is_large and line.bold and is_short:
            return HeadingCandidate(
                level=self.size_rank.get(round(line.size, 1), 1), title=title
            )
        return None


class SectionStack:
    def __init__(self):
        self._stack: list[tuple[int, str]] = []

    def push(self, level: int, title: str) -> None:
        while self._stack and self._stack[-1][0] >= level:
            self._stack.pop()
        self._stack.append((level, title))

    @property
    def path(self) -> str:
        return " > ".join(title for _, title in self._stack)

    @property
    def current_title(self) -> str:
        return self._stack[-1][1] if self._stack else ""


def _word_center_in_bbox(word: dict, bbox: tuple[float, float, float, float]) -> bool:
    x0, top, x1, bottom = bbox
    cx = (word["x0"] + word["x1"]) / 2
    cy = (word["top"] + word["bottom"]) / 2
    return x0 <= cx <= x1 and top <= cy <= bottom


def _rows_to_text(rows) -> str:
    lines = []
    for row in rows:
        cells = [(cell or "").strip().replace("\n", " ") for cell in row]
        if any(cells):
            lines.append(" | ".join(cells))
    return "\n".join(lines)


def _digit_collapsed(text: str) -> str:
    # Digit-insensitive signature: a running header repeating "Page 10 of 35"
    # would otherwise defeat exact-text repeat detection.
    return _DIGITS_RE.sub("#", text)


def _page_words(page) -> list[dict]:
    """Words with font size and boldness. `get_text("words")` is faster but
    drops font attributes, which heading detection needs, so spans are split
    on whitespace instead and each word inherits its span's font.
    """
    words: list[dict] = []
    for block in page.get_text("dict").get("blocks", []):
        if block.get("type") != 0:  # 0 = text, 1 = image
            continue
        for line in block.get("lines", []):
            for span in line.get("spans", []):
                span_text = span.get("text") or ""
                if not span_text.strip():
                    continue
                x0, top, x1, bottom = span["bbox"]
                is_bold = bool(span.get("flags", 0) & _BOLD_FLAG) or "bold" in str(
                    span.get("font", "")
                ).lower()
                fontname = span.get("font", "") + ("-Bold" if is_bold else "")

                # Span bboxes cover the whole span, so per-word x extents are
                # apportioned by character count - approximate is fine, they
                # only order words within a line and test table containment.
                char_width = (x1 - x0) / max(len(span_text), 1)
                cursor = 0
                for token in span_text.split():
                    start = span_text.index(token, cursor)
                    cursor = start + len(token)
                    words.append(
                        {
                            "text": token,
                            "x0": x0 + start * char_width,
                            "x1": x0 + cursor * char_width,
                            "top": top,
                            "bottom": bottom,
                            "size": span.get("size", 0.0),
                            "fontname": fontname,
                        }
                    )
    return words


def _looks_tabular(rows) -> bool:
    """Rejects false-positive tables: PyMuPDF's finder keys off ruling lines,
    so a decorative box comes back as a 2x2 table of sentence fragments.
    """
    cells = [[(c or "").strip() for c in row] for row in rows]
    if max((len(row) for row in cells), default=0) < 2:
        return False
    filled = sum(1 for row in cells for c in row if c)
    multi_column_rows = sum(1 for row in cells if sum(1 for c in row if c) >= 2)
    return multi_column_rows >= _MIN_TABLE_ROWS and filled >= 6


def _page_tables(page) -> list[dict]:
    """{bbox, text} per table, so a table can be emitted as its own unit and
    have its words excluded from the prose stream.
    """
    tables: list[dict] = []
    try:
        found = page.find_tables()
    except Exception:  # noqa: BLE001 - table finding is best-effort
        return tables

    for table in found.tables:
        try:
            data = table.extract()
        except Exception:  # noqa: BLE001
            continue
        if len(data) < _MIN_TABLE_ROWS or not _looks_tabular(data):
            continue
        text = _rows_to_text(data)
        if text.strip():
            tables.append({"bbox": tuple(table.bbox), "text": text})
    return tables


class _SectionUnitBuilder:
    """Walks a PDF's (line, table) events in reading order, flushing prose
    into per-section units and emitting tables standalone, tagged with the
    section active when they appear.
    """

    def __init__(
        self,
        classifier: HeadingClassifier,
        noisy_table_signatures: set[str],
        noisy_line_signatures: set[str],
    ):
        self._classifier = classifier
        self._noisy_table_signatures = noisy_table_signatures
        self._noisy_line_signatures = noisy_line_signatures
        self.units: list[ExtractedUnit] = []
        self._stack = SectionStack()
        self._buffer: list[str] = []
        self._buffer_page_start: int | None = None
        self._last_body_line_text = ""
        self._table_counter = 0

    def process_page(self, page_index: int, lines: list[Line], tables: list[dict]) -> None:
        events: list[tuple[float, str, object]] = [(line.top, "line", line) for line in lines]
        events += [(t["bbox"][1], "table", t) for t in tables]
        events.sort(key=lambda e: e[0])

        for _, kind, obj in events:
            if kind == "line":
                self._handle_line(page_index, obj)
            else:
                self._handle_table(page_index, obj)

    def _handle_line(self, page_index: int, line: Line) -> None:
        if _digit_collapsed(line.text.strip()) in self._noisy_line_signatures:
            return
        heading = self._classifier.classify(line)
        if heading:
            self.flush()
            self._stack.push(heading.level, heading.title)
            # Don't let the next table inherit a title from a prior section.
            self._last_body_line_text = ""
            return
        if self._buffer_page_start is None:
            self._buffer_page_start = page_index
        self._buffer.append(line.text)
        self._last_body_line_text = line.text

    def _handle_table(self, page_index: int, table: dict) -> None:
        if _digit_collapsed(table["text"]) in self._noisy_table_signatures:
            return
        self.flush()
        self._table_counter += 1
        self.units.append(
            ExtractedUnit(
                text=table["text"],
                unit_type="table",
                locator={
                    "page": page_index,
                    "section_path": self._stack.path,
                    "table_title": self._last_body_line_text or self._stack.current_title,
                    "table": self._table_counter,
                },
            )
        )

    def flush(self) -> None:
        if self._buffer:
            text = "\n".join(self._buffer)
            if text.strip():
                self.units.append(
                    ExtractedUnit(
                        text=text,
                        unit_type="section_text",
                        locator={
                            "page": self._buffer_page_start,
                            "section_path": self._stack.path,
                            "section_title": self._stack.current_title,
                        },
                    )
                )
            self._buffer.clear()
        self._buffer_page_start = None


class PDFExtractor:
    """Structure-aware PDF extraction: builds a heading hierarchy (numbering
    patterns + font-size/bold tiers) while walking the document, emits prose
    as per-section units that never span two sections, and tables as their
    own units. Running headers/footers and repeated admin tables are dropped.
    """

    def __init__(
        self,
        heading_size_ratio: float = _HEADING_SIZE_RATIO_DEFAULT,
        table_repeat_threshold: float = _TABLE_REPEAT_THRESHOLD_DEFAULT,
    ):
        self.heading_size_ratio = heading_size_ratio
        self.table_repeat_threshold = table_repeat_threshold

    def extract(self, file_path: Path) -> list[ExtractedUnit]:
        pages_data = self._read_pages(file_path)
        if not pages_data:
            return []

        builder = _SectionUnitBuilder(
            self._build_classifier(pages_data),
            self._detect_repeated_tables(pages_data),
            self._detect_repeated_lines(pages_data),
        )
        for page_index, lines, tables in pages_data:
            builder.process_page(page_index, lines, tables)
        builder.flush()
        return builder.units

    def _read_pages(self, file_path: Path) -> list[tuple[int, list[Line], list[dict]]]:
        pages_data: list[tuple[int, list[Line], list[dict]]] = []
        with pymupdf.open(file_path) as doc:
            for page_index, page in enumerate(doc, start=1):
                tables = _page_tables(page)
                words = [
                    w
                    for w in _page_words(page)
                    if not any(_word_center_in_bbox(w, t["bbox"]) for t in tables)
                ]
                pages_data.append(
                    (page_index, group_words_into_lines(words, page_index), tables)
                )
        return pages_data

    def _build_classifier(self, pages_data) -> HeadingClassifier:
        body_size = compute_body_size(
            [line.size for _, lines, _ in pages_data for line in lines]
        )
        candidate_sizes = [
            line.size
            for _, lines, _ in pages_data
            for line in lines
            if line.bold
            and body_size
            and line.size > body_size * self.heading_size_ratio
            and len(line.text.split()) <= _FALLBACK_HEADING_MAX_WORDS
        ]
        return HeadingClassifier(
            body_size, build_size_rank(candidate_sizes), self.heading_size_ratio
        )

    def _detect_repeated_tables(self, pages_data) -> set[str]:
        """Tables whose text repeats across most pages are running boilerplate."""
        signature_pages: dict[str, set[int]] = {}
        for page_index, _, tables in pages_data:
            for t in tables:
                signature_pages.setdefault(_digit_collapsed(t["text"]), set()).add(page_index)

        threshold = max(2, round(len(pages_data) * self.table_repeat_threshold))
        return {sig for sig, pages in signature_pages.items() if len(pages) >= threshold}

    def _detect_repeated_lines(self, pages_data) -> set[str]:
        """Lines repeating across most pages are headers/footers. Detected per
        line rather than per page blob, so a footer can never leak into a
        section's text or be mistaken for a heading.
        """
        if len(pages_data) < 3:
            return set()

        line_pages: dict[str, set[int]] = {}
        for page_index, lines, _ in pages_data:
            texts = {line.text.strip() for line in lines if line.text.strip()}
            for sig in {_digit_collapsed(text) for text in texts}:
                line_pages.setdefault(sig, set()).add(page_index)

        threshold = max(2, round(len(pages_data) * _LINE_REPEAT_THRESHOLD))
        return {text for text, pages in line_pages.items() if len(pages) >= threshold}


# --- XLSX ------------------------------------------------------------------

_TEST_CASE_ID_ALIASES = {"test case id", "test_case_id", "testcaseid", "tc id", "tcid", "case id"}
_REQUIREMENT_ID_ALIASES = {"requirement", "requirement id", "requirement_id", "req id", "reqid"}
_STEPS_ALIASES = {"steps", "test steps", "test_steps", "step"}
_EXPECTED_RESULT_ALIASES = {"expected result", "expected_result", "expected", "result"}
_SEQUENCE_ALIASES = {"s_no", "s.no", "sno", "sr no", "sl no", "seq", "id", "#"}


def _cell_str(value) -> str:
    return "" if value is None else str(value).strip()


def _match_column(header: list[str], aliases: set[str]) -> int | None:
    for index, name in enumerate(header):
        if name.strip().lower() in aliases:
            return index
    return None


class XLSXExtractor:
    """One unit per non-empty data row (unit_type="test_case"), with the
    sheet's header prefixed as "Header: value" pairs so a row stays
    self-describing once separated from its sheet. Known columns
    (requirement id, test case id, ...) are also pulled into `extra`.
    """

    def extract(self, file_path: Path) -> list[ExtractedUnit]:
        units: list[ExtractedUnit] = []
        workbook = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
        try:
            for sheet in workbook.worksheets:
                units.extend(self._extract_sheet(sheet))
        finally:
            workbook.close()
        return units

    def _extract_sheet(self, sheet) -> list[ExtractedUnit]:
        rows = sheet.iter_rows(values_only=True)
        try:
            header = [_cell_str(v) for v in next(rows)]
        except StopIteration:
            return []

        columns = {
            "requirement_id": _match_column(header, _REQUIREMENT_ID_ALIASES),
            "test_case_id": _match_column(header, _TEST_CASE_ID_ALIASES),
            "steps": _match_column(header, _STEPS_ALIASES),
            "expected_result": _match_column(header, _EXPECTED_RESULT_ALIASES),
            "sequence": _match_column(header, _SEQUENCE_ALIASES),
        }

        units: list[ExtractedUnit] = []
        for row_index, row in enumerate(rows, start=2):
            unit = self._build_row_unit(
                sheet.title, row_index, header, [_cell_str(v) for v in row], columns
            )
            if unit:
                units.append(unit)
        return units

    @staticmethod
    def _build_row_unit(
        sheet_title: str,
        row_index: int,
        header: list[str],
        values: list[str],
        columns: dict[str, int | None],
    ) -> ExtractedUnit | None:
        if not any(values):
            return None

        pairs = [f"{h}: {v}" for h, v in zip(header, values) if h and v]
        if not pairs:
            return None

        def get(col_name: str) -> str | None:
            index = columns[col_name]
            return values[index] if index is not None and values[index] else None

        requirement_id = get("requirement_id")
        test_case_id = get("test_case_id")
        sequence = get("sequence")
        if not test_case_id and requirement_id and sequence:
            test_case_id = f"{requirement_id}_{sequence}"

        return ExtractedUnit(
            text=" | ".join(pairs),
            unit_type="test_case",
            locator={"sheet": sheet_title, "row": row_index},
            extra={
                "header": header,
                "requirement_id": requirement_id,
                "test_case_id": test_case_id,
                "steps": get("steps"),
                "expected_result": get("expected_result"),
            },
        )


# --- Python ----------------------------------------------------------------

_METHOD_TYPES = (ast.FunctionDef, ast.AsyncFunctionDef)


def _node_source(lines: list[str], node: ast.AST) -> str:
    """Source for a def/class node including its decorators, which
    ast.get_source_segment excludes.
    """
    start_line = node.lineno
    for decorator in getattr(node, "decorator_list", []):
        start_line = min(start_line, decorator.lineno)
    return "\n".join(lines[start_line - 1 : node.end_lineno])


def _is_docstring_stmt(node: ast.stmt) -> bool:
    return (
        isinstance(node, ast.Expr)
        and isinstance(node.value, ast.Constant)
        and isinstance(node.value.value, str)
    )


def _class_header_source(lines: list[str], node: ast.ClassDef) -> str:
    """The `class Foo(Base):` signature plus docstring only - used when the
    class's methods are chunked separately.
    """
    start_line = node.lineno
    for decorator in getattr(node, "decorator_list", []):
        start_line = min(start_line, decorator.lineno)

    end_line = node.lineno
    if node.body:
        first = node.body[0]
        end_line = max(start_line, first.lineno - 1)
        if _is_docstring_stmt(first):
            end_line = first.end_lineno

    return "\n".join(lines[start_line - 1 : end_line])


class PythonExtractor:
    """Module -> class -> method structure, so a chunk never splits a method.

    A module docstring is its own unit; a class with methods becomes a
    compact class-header unit plus one unit per method; a class without
    methods stays whole. Falls back to whole-file text on a syntax error.
    """

    def extract(self, file_path: Path) -> list[ExtractedUnit]:
        source = file_path.read_text(encoding="utf-8", errors="replace")

        try:
            tree = ast.parse(source, filename=str(file_path))
        except SyntaxError:
            return [
                ExtractedUnit(
                    text=source, unit_type="raw_python", locator={"reason": "syntax_error"}
                )
            ]

        units: list[ExtractedUnit] = []
        lines = source.splitlines()

        module_docstring = ast.get_docstring(tree)
        if module_docstring:
            units.append(
                ExtractedUnit(
                    text=module_docstring, unit_type="module_docstring", locator={"lineno": 1}
                )
            )

        for node in tree.body:
            if isinstance(node, ast.ClassDef):
                units.extend(self._extract_class(lines, node))
            elif isinstance(node, _METHOD_TYPES):
                units.extend(self._extract_function(lines, node))

        if not units:
            # Only top-level statements - keep the whole file.
            units.append(
                ExtractedUnit(
                    text=source, unit_type="raw_python", locator={"reason": "no_top_level_defs"}
                )
            )
        return units

    def _extract_class(self, lines: list[str], node: ast.ClassDef) -> list[ExtractedUnit]:
        methods = [n for n in node.body if isinstance(n, _METHOD_TYPES)]
        text = _class_header_source(lines, node) if methods else _node_source(lines, node)

        units: list[ExtractedUnit] = []
        if text.strip():
            units.append(
                ExtractedUnit(
                    text=text,
                    unit_type="class",
                    locator={"name": node.name, "lineno": node.lineno},
                    extra={"docstring": ast.get_docstring(node) or ""},
                )
            )

        for method in methods:
            method_text = _node_source(lines, method)
            if not method_text.strip():
                continue
            units.append(
                ExtractedUnit(
                    text=method_text,
                    unit_type="function",
                    locator={
                        "name": method.name,
                        "lineno": method.lineno,
                        "class_name": node.name,
                    },
                    extra={"docstring": ast.get_docstring(method) or ""},
                )
            )
        return units

    def _extract_function(self, lines: list[str], node) -> list[ExtractedUnit]:
        text = _node_source(lines, node)
        if not text.strip():
            return []
        return [
            ExtractedUnit(
                text=text,
                unit_type="function",
                locator={"name": node.name, "lineno": node.lineno},
                extra={"docstring": ast.get_docstring(node) or ""},
            )
        ]


# --- Plain text ------------------------------------------------------------

_PYTHON_MARKERS = ("import ", "def ", "class ", "elif ", "self.")
_MIN_PYTHON_MARKERS = 2


def looks_like_python(text: str) -> bool:
    """Whether a .txt file is really Python source.

    The reference automation scripts here are Python saved as .txt. Treated
    as prose they get sentence-split and space-collapsed, destroying the
    indentation that makes them useful - so content decides the parser, not
    the extension. Markers are checked before ast.parse because a single
    bare word is also valid Python.
    """
    if sum(marker in text for marker in _PYTHON_MARKERS) < _MIN_PYTHON_MARKERS:
        return False
    try:
        ast.parse(text)
    except (SyntaxError, ValueError):
        return False
    return True


class TextExtractor:
    """Prose becomes one text unit; a .txt that is really Python is handed to
    PythonExtractor so it keeps its structure and layout.
    """

    def __init__(self):
        self._python_extractor = PythonExtractor()

    def extract(self, file_path: Path) -> list[ExtractedUnit]:
        text = file_path.read_text(encoding="utf-8", errors="replace")
        if not text.strip():
            return []
        if looks_like_python(text):
            return self._python_extractor.extract(file_path)
        return [ExtractedUnit(text=text, unit_type="text", locator={})]


# --- HTML track data -------------------------------------------------------

# Reports identify their subdivision three ways: a body line, the document
# title, and a quoted display name. The numeric ID is what the automation
# scripts pass as subdivID, so it is read first.
_SUBDIVISION_ID_RE = re.compile(r"subdivision\s+(\d{3,6})\b", re.IGNORECASE)
_TITLE_ID_RE = re.compile(r"<title>[^<]*?\.?(\d{4,6})\.\d+", re.IGNORECASE)
_SUBDIVISION_NAME_RE = re.compile(r'Subdivision\s*"([^"]+)"')
_STEM_ID_RE = re.compile(r"(\d{3,6})")
_WHITESPACE_RE = re.compile(r"\s+")


class _TrackReportParser(HTMLParser):
    """Parses a Wabtec 'HTML Track by Group' report: `<a name="Section">`
    followed by a table whose first row is column labels and the rest
    track-feature records.

    Some cells render their content as a nested <table>. `_table_depth`
    tracks that nesting so only the outermost table drives row extraction -
    otherwise a nested table desyncs the row state and silently drops every
    real row that follows.
    """

    def __init__(self):
        super().__init__()
        self.sections: list[tuple[str, list[list[str]]]] = []
        self._pending_section_name: str | None = None
        self._table_depth = 0
        self._in_row = False
        self._in_cell = False
        self._current_table: list[list[str]] = []
        self._current_row: list[str] = []
        self._current_cell_parts: list[str] = []

    def handle_starttag(self, tag, attrs):
        if tag == "a":
            name = dict(attrs).get("name")
            if name:
                self._pending_section_name = name
            return

        if tag == "table":
            self._table_depth += 1
            if self._table_depth == 1:
                self._current_table = []
            return

        if tag == "tr":
            if self._table_depth == 1:
                self._in_row = True
                self._current_row = []
            elif self._table_depth > 1 and self._in_cell and self._has_cell_content():
                self._current_cell_parts.append("; ")
            return

        if tag in ("td", "th"):
            if self._table_depth == 1 and self._in_row:
                self._in_cell = True
                self._current_cell_parts = []
            elif self._table_depth > 1 and self._in_cell and self._has_cell_content():
                self._current_cell_parts.append(" ")
            return

        if tag == "br" and self._in_cell:
            self._current_cell_parts.append("; ")

    def handle_endtag(self, tag):
        if tag in ("td", "th"):
            if self._table_depth == 1 and self._in_cell:
                text = _WHITESPACE_RE.sub(" ", "".join(self._current_cell_parts)).strip()
                self._current_row.append(text)
                self._in_cell = False
            return

        if tag == "tr":
            if self._table_depth == 1 and self._in_row:
                if any(self._current_row):
                    self._current_table.append(self._current_row)
                self._in_row = False
            return

        if tag == "table":
            if self._table_depth == 1:
                if self._pending_section_name and len(self._current_table) >= _MIN_TABLE_ROWS:
                    self.sections.append((self._pending_section_name, self._current_table))
                self._pending_section_name = None
                self._current_table = []
            if self._table_depth > 0:
                self._table_depth -= 1

    def handle_data(self, data):
        if self._in_cell:
            self._current_cell_parts.append(data)

    def _has_cell_content(self) -> bool:
        # Whitespace-only nodes between a cell's open tag and a nested
        # table's first row would otherwise produce a leading separator.
        return bool("".join(self._current_cell_parts).strip())


def normalize_subdivision(value: str) -> str:
    """"8880", "08880" and " 8880 " all name the same subdivision; pad to five."""
    digits = "".join(c for c in str(value) if c.isdigit())
    return digits.zfill(5) if digits else str(value).strip()


def _subdivision_id(raw_html: str, file_path: Path) -> str:
    """The numeric subdivision ID, tried in order of reliability: the body
    line that states it, the document title, then the filename - any one of
    them can be missing or reformatted by whoever exported the report.
    """
    for pattern, source in (
        (_SUBDIVISION_ID_RE, raw_html),
        (_TITLE_ID_RE, raw_html),
        (_STEM_ID_RE, file_path.stem),
    ):
        match = pattern.search(source)
        if match:
            return normalize_subdivision(match.group(1))
    return file_path.stem


class HTMLTrackDataExtractor:
    """Wabtec subdivision reports: subdivision -> named section (Blocks,
    Switches, Signals, Speed Restrictions) -> track-feature records. Each
    section becomes one unit_type="table" unit, the same shape as a PDF
    table, so it needs no special-casing downstream.
    """

    def extract(self, file_path: Path) -> list[ExtractedUnit]:
        raw_html = file_path.read_text(encoding="utf-8", errors="replace")

        subdivision = _subdivision_id(raw_html, file_path)
        name_match = _SUBDIVISION_NAME_RE.search(raw_html)
        subdivision_name = name_match.group(1).strip() if name_match else ""
        label = f"{subdivision} {subdivision_name}".strip()

        parser = _TrackReportParser()
        parser.feed(raw_html)

        units: list[ExtractedUnit] = []
        for index, (section_name, rows) in enumerate(parser.sections, start=1):
            text = _rows_to_text(rows)
            if not text.strip():
                continue
            units.append(
                ExtractedUnit(
                    text=text,
                    unit_type="table",
                    locator={
                        "section_path": f"{label} > {section_name}",
                        "table_title": section_name,
                        "table": index,
                        # Kept for citation: retrieval searches every subdivision.
                        "subdivision": subdivision,
                        "subdivision_name": subdivision_name or None,
                    },
                )
            )
        return units


_python_extractor = PythonExtractor()
_xlsx_extractor = XLSXExtractor()
_html_extractor = HTMLTrackDataExtractor()

EXTRACTORS_BY_SUFFIX: dict[str, Extractor] = {
    ".pdf": PDFExtractor(),
    ".xlsx": _xlsx_extractor,
    ".xlsm": _xlsx_extractor,
    ".py": _python_extractor,
    ".pyi": _python_extractor,
    ".txt": TextExtractor(),
    ".html": _html_extractor,
    ".htm": _html_extractor,
}
